"""Robô TOTALE: descoberta, leitura de CSV/Excel e sincronismo na Central.

O motor, o ciclo headless e a interface Streamlit vivem neste único arquivo.
A dependência de Streamlit é opcional para a execução via CLI e para os testes
de leitura; o portal importa somente :func:`renderizar_robo_local`.

Arquivos aceitos na pasta monitorada
-----------------------------------
``Atividades-*.csv``, ``Atividades-*.xlsx``, ``Atividades-*.xlsm`` e
``Atividades-*.xls`` — somente arquivos cujo nome inicia-se com
``Atividades-``. Quando há mais de um arquivo, o pipeline usa **somente o mais
recente** da pasta inteira; os demais aparecem na fila apenas como histórico.
"""

from __future__ import annotations

import csv
import fnmatch
import logging
import math
import os
import sys
import tempfile
import time
import warnings
import zipfile
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from datetime import datetime
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Literal, cast

import pandas as pd

try:
    from pandas.errors import OptionError as PandasOptionError
except ImportError:  # pandas < 1.0
    PandasOptionError = KeyError  # type: ignore[misc, assignment]

logger = logging.getLogger(__name__)

# ── Constantes ──────────────────────────────────────────────────────
EngineExcel = Literal["openpyxl", "calamine", "xlrd"]
ModoLeitura = Literal["ultimo", "ultimo_por_base", "todos"]
# caminho, mtime em nanossegundos, tamanho em bytes (assinatura sem perda de precisão)
Candidato = tuple[str, int, int]

PADROES_TOTALE: tuple[str, ...] = (
    "Atividades-*.csv",
    "Atividades-*.xlsx",
    "Atividades-*.xlsm",
    "Atividades-*.xls",
)
EXTENSOES_VALIDAS: tuple[str, ...] = (".csv", ".xlsx", ".xlsm", ".xls")
EXTENSOES_TEMP: tuple[str, ...] = (
    ".tmp",
    ".crdownload",
    ".partial",
    ".download",
    ".part",
)

MAPA_BASES_ARQUIVO: dict[str, str] = {
    "ABCDM": "NET-ABCDM",
    "SPO": "NET-LESTE",
    "GRS": "NET-GUARULHOS",
}

_VALORES_AUSENTES: frozenset[str] = frozenset(
    {
        "",
        "nan",
        "none",
        "na",
        "n/a",
        "<na>",
        "nat",
        "null",
        "não informado",
        "nao informado",
        "-",
    }
)
_TAMANHO_MINIMO_XLSX: int = 2_000
_TAMANHO_MAXIMO_PADRAO: int = 80 * 1024 * 1024


class ArquivoBloqueadoError(Exception):
    """Arquivo em gravação ou sincronização. O próximo ciclo tenta de novo."""


class _BytesNomeados(BytesIO):
    """BytesIO com .name, para o pandas não inferir a extensão como zip."""

    def __init__(self, raw: bytes, nome: str) -> None:
        super().__init__(raw)
        self.name = nome


@dataclass(frozen=True)
class LeituraArquivo:
    df: pd.DataFrame | None
    erro: str | None
    motor: str | None


@dataclass(frozen=True)
class ResultadoCiclo:
    df: pd.DataFrame | None
    bases: tuple[str, ...]
    arquivos: tuple[str, ...]
    erros: tuple[str, ...]
    bloqueio: str | None
    motores: tuple[str, ...]
    duplicatas_removidas: int


# ── Texto e colunas ─────────────────────────────────────────────────
def _celula_para_texto(valor: object) -> str:
    if valor is None or valor is pd.NA:
        return ""
    if isinstance(valor, bool):
        return "True" if valor else "False"
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if math.isnan(valor):
            return ""
        if valor.is_integer():
            return str(int(valor))
        return str(valor).strip()
    texto = valor.strip() if isinstance(valor, str) else str(valor).strip()
    if texto.lower() in {"nat", "<na>", "nan", "none"}:
        return ""
    corpo = texto[:-2] if texto.endswith(".0") else ""
    if corpo and corpo.replace("-", "", 1).isdigit():
        return corpo
    return texto


def _nomes_colunas_unicos(nomes: Iterable[object]) -> list[str]:
    limpos: list[str] = []
    for i, bruto in enumerate(nomes):
        nome = (
            "" if bruto is None else str(bruto).replace("\ufeff", "").replace("\xa0", " ").strip()
        )
        limpos.append(f"col_{i}" if not nome or nome.lower().startswith("unnamed") else nome)

    # Reservar as grafias originais impede que uma duplicata de "A" vire "A.1"
    # quando já existe uma coluna chamada "A.1" no arquivo.
    reservados = set(limpos)
    usados: set[str] = set()
    contadores: dict[str, int] = {}
    saida: list[str] = []
    for nome in limpos:
        unico = nome
        if unico in usados:
            sufixo = contadores.get(nome, 0) + 1
            while f"{nome}.{sufixo}" in reservados or f"{nome}.{sufixo}" in usados:
                sufixo += 1
            contadores[nome] = sufixo
            unico = f"{nome}.{sufixo}"
        usados.add(unico)
        saida.append(unico)
    return saida


def _padronizar_coluna_texto(serie: pd.Series) -> pd.Series:
    """Converte uma coluna para texto limpo, de forma **vetorial**.

    Equivalente a aplicar :func:`_celula_para_texto` célula a célula, porém sem
    callback Python por valor. Em bases reais (120 mil linhas × 18 colunas) a
    padronização caiu de ~0,82 s para ~0,13 s (**6× mais rápida**) com saída
    idêntica — ver ``tests/test_robo_local.py``.
    """
    if pd.api.types.is_bool_dtype(serie):
        return serie.map({True: "True", False: "False"}).astype("string")

    if pd.api.types.is_datetime64_any_dtype(serie):
        # str(Timestamp) → "AAAA-MM-DD HH:MM:SS"; mantém o formato anterior.
        return serie.dt.strftime("%Y-%m-%d %H:%M:%S").astype("string")

    if pd.api.types.is_numeric_dtype(serie):
        numerica = pd.to_numeric(serie, errors="coerce")
        texto = numerica.astype("string")
        inteiros = numerica.notna() & (numerica == numerica.round())
        if inteiros.any():
            # Não converte a coluna inteira para Int64: isso estoura para valores
            # integrais > 2**63-1, apesar de haver floats perfeitamente válidos.
            com_decimal_zero = inteiros & texto.str.fullmatch(r"-?\d+\.0+", na=False)
            texto = texto.mask(com_decimal_zero, texto.str.replace(r"\.0+$", "", regex=True))
            notacao_cientifica = inteiros & texto.str.contains(r"[eE]", na=False)
            if notacao_cientifica.any():
                texto.loc[notacao_cientifica] = numerica.loc[notacao_cientifica].map(
                    _celula_para_texto
                )
        # "-0.0" também é zero para ``_celula_para_texto``; alinhe as duas rotas.
        texto = texto.mask(numerica == 0, "0")
        # Decimais em notação científica seguem o ``str(float)`` da rota célula a
        # célula; o dtype string do pandas usa outra precisão em alguns casos.
        nao_inteiros = numerica.notna() & ~inteiros
        cientificos = nao_inteiros & texto.str.contains(r"[eE]", na=False)
        if cientificos.any():
            texto.loc[cientificos] = numerica.loc[cientificos].map(_celula_para_texto)
        return texto

    texto = serie.astype("string").str.strip()
    com_sufixo = texto.str.endswith(".0", na=False).fillna(False)
    if com_sufixo.any():
        corpo = texto.str[:-2]
        digitos = corpo.str.fullmatch(r"-?\d+", na=False).fillna(False)
        texto = texto.mask(com_sufixo & digitos, corpo)
    return texto


def _padronizar_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty and len(df.columns) == 0:
        return df.copy()
    saida = df.copy()
    saida.columns = pd.Index(_nomes_colunas_unicos(list(saida.columns)))
    for col in saida.columns:
        texto = _padronizar_coluna_texto(saida[col]).str.strip()
        saida[col] = texto.mask(texto.str.lower().isin(_VALORES_AUSENTES))
    return saida.dropna(how="all")


def _cabecalho_fraco(colunas: Sequence[object]) -> bool:
    nomes = [str(c).strip() for c in colunas]
    if not nomes:
        return True
    ruins = sum(
        1 for n in nomes if not n or n.lower().startswith("unnamed") or n.lower().startswith("col_")
    )
    return ruins / len(nomes) >= 0.5


def _alinhar_colunas_esperadas(
    df: pd.DataFrame,
    esperadas: Sequence[str] | None,
) -> pd.DataFrame:
    if not esperadas:
        return df
    por_chave = {str(c).strip().casefold(): str(c) for c in df.columns}
    renomear: dict[str, str] = {}
    for esperada in esperadas:
        limpa = str(esperada).strip()
        if limpa in df.columns:
            continue
        atual = por_chave.get(limpa.casefold())
        if atual and atual not in renomear:
            renomear[atual] = limpa
    if not renomear:
        return df
    return df.rename(columns=renomear)


def _chaves_colunas(df: pd.DataFrame) -> set[str]:
    return {str(c).strip().casefold() for c in df.columns}


def _contem_colunas(df: pd.DataFrame, esperadas: Sequence[str] | None) -> bool:
    if not esperadas:
        return True
    chaves = _chaves_colunas(df)
    return all(str(esperada).strip().casefold() in chaves for esperada in esperadas)


def _faltam_colunas(df: pd.DataFrame, esperadas: Sequence[str] | None) -> list[str]:
    if not esperadas:
        return []
    chaves = _chaves_colunas(df)
    return [c for c in esperadas if str(c).strip().casefold() not in chaves]


# ── Detecção de base ────────────────────────────────────────────────
def detectar_marcador_arquivo(caminho_ou_nome: str | Path | None) -> str | None:
    if not caminho_ou_nome:
        return None
    nome = Path(str(caminho_ou_nome)).name.upper()
    for marcador in sorted(MAPA_BASES_ARQUIVO, key=len, reverse=True):
        if marcador in nome:
            return marcador
    return None


def detectar_base_arquivo(caminho_ou_nome: str | Path | None) -> str | None:
    marcador = detectar_marcador_arquivo(caminho_ou_nome)
    if marcador is None:
        return None
    return MAPA_BASES_ARQUIVO.get(marcador)


def _canonizar_coluna_base(df: pd.DataFrame) -> pd.DataFrame:
    variantes = [col for col in df.columns if str(col).casefold() == "base" and col != "BASE"]
    if not variantes:
        return df
    saida = df.copy()
    if "BASE" not in saida.columns:
        saida = saida.rename(columns={variantes[0]: "BASE"})
        variantes = variantes[1:]
    for col in variantes:
        texto = saida["BASE"].map(_celula_para_texto).str.strip().str.lower()
        vazios = texto.isin(_VALORES_AUSENTES)
        saida.loc[vazios, "BASE"] = saida.loc[vazios, col]
        saida = saida.drop(columns=[col])
    return saida


def _aplicar_base_no_dataframe(df: pd.DataFrame, base: str) -> pd.DataFrame:
    saida = _canonizar_coluna_base(df)
    if "BASE" not in saida.columns:
        saida["BASE"] = base
        return saida
    texto = saida["BASE"].map(_celula_para_texto).str.strip().str.lower()
    vazios = texto.isin(_VALORES_AUSENTES)
    saida.loc[vazios, "BASE"] = base
    return saida


def _unificar_nomes_colunas(frames: list[pd.DataFrame]) -> list[pd.DataFrame]:
    """Uma grafia por coluna na junção. O primeiro arquivo define o nome."""
    canonico: dict[str, str] = {}
    unificados: list[pd.DataFrame] = []
    for df in frames:
        renomear: dict[str, str] = {}
        for col in df.columns:
            chave = str(col).casefold()
            nome = canonico.get(chave)
            if nome is None:
                canonico[chave] = str(col)
            elif nome != col:
                renomear[str(col)] = nome
        unificados.append(df.rename(columns=renomear) if renomear else df)
    return unificados


# ── Descoberta e Metadados de Arquivos ──────────────────────────────
def obter_pasta_robo_padrao() -> str:
    """Pasta padrão do robô: C:\\Users\\<usuario>\\downloads."""
    robo = Path.home() / "downloads"
    robo.mkdir(parents=True, exist_ok=True)
    return str(robo.resolve())


def formatar_tamanho(bytes_: int) -> str:
    if bytes_ < 1024:
        return f"{bytes_} B"
    if bytes_ < 1024 * 1024:
        return f"{bytes_ / 1024:.1f} KB"
    return f"{bytes_ / (1024 * 1024):.2f} MB"


def obter_metadados_arquivo(caminho: str | None) -> dict[str, str]:
    vazio = {
        "nome": Path(caminho).name if caminho else "Desconhecido",
        "ext": Path(caminho).suffix.lower() if caminho else "",
        "tamanho": "—",
        "modificado": "—",
        "caminho_completo": caminho or "",
    }
    if not caminho or not os.path.exists(caminho):
        return vazio
    try:
        p = Path(caminho)
        stat = p.stat()
        return {
            "nome": p.name,
            "ext": p.suffix.lower(),
            "tamanho": formatar_tamanho(stat.st_size),
            "modificado": datetime.fromtimestamp(stat.st_mtime).strftime("%d/%m/%Y %H:%M:%S"),
            "caminho_completo": str(p.resolve()),
        }
    except OSError:
        return vazio


def _stat_basico(caminho: str) -> tuple[int | None, int | None]:
    try:
        st_ = os.stat(caminho)
        return int(st_.st_mtime_ns), int(st_.st_size)
    except OSError:
        return None, None


def _arquivo_ignorado(nome: str) -> bool:
    baixo = nome.lower()
    if baixo.startswith(("~$", ".")):
        return True
    return baixo.endswith(EXTENSOES_TEMP)


def _listar_arquivos(pasta: Path, recursivo: bool) -> list[Path]:
    try:
        if recursivo:
            return [p for p in pasta.rglob("*") if p.is_file()]
        return [p for p in pasta.iterdir() if p.is_file()]
    except OSError as e:
        logger.warning("Não foi possível listar %s: %s", pasta, e)
        return []


def _casa_padrao(nome: str, padroes: Sequence[str]) -> bool:
    baixo = nome.lower()
    return any(fnmatch.fnmatch(baixo, p.lower()) for p in padroes)


def _listar_candidatos_com_stat(
    pasta: str,
    padroes: tuple[str, ...] = PADROES_TOTALE,
    recursivo: bool = False,
) -> list[Candidato]:
    if not str(pasta).strip():
        return []
    p = Path(pasta)
    if not p.is_dir():
        return []

    arquivos = [arq for arq in _listar_arquivos(p, recursivo) if not _arquivo_ignorado(arq.name)]
    encontrados = [
        arq
        for arq in arquivos
        if arq.suffix.lower() in EXTENSOES_VALIDAS
        and arq.name.lower().startswith("atividades-")
        and _casa_padrao(arq.name, padroes)
    ]

    candidatos: list[Candidato] = []
    for arq in encontrados:
        mtime, size = _stat_basico(str(arq))
        if mtime is None or size is None or size <= 0:
            continue
        candidatos.append((str(arq), mtime, size))
    candidatos.sort(key=lambda item: item[1], reverse=True)
    return candidatos


def buscar_arquivo_mais_recente(
    pasta: str,
    padroes: tuple[str, ...] = PADROES_TOTALE,
) -> tuple[str | None, float | None, int | None]:
    """Retorna o arquivo mais novo; mantém mtime em segundos na API pública."""
    candidatos = _listar_candidatos_com_stat(pasta, padroes)
    if not candidatos:
        return None, None, None
    caminho, mtime_ns, size = candidatos[0]
    return caminho, mtime_ns / 1_000_000_000, size


def mapear_arquivos_por_base(pasta: str, recursivo: bool = False) -> dict[str, Candidato]:
    por_base: dict[str, Candidato] = {}
    for item in _listar_candidatos_com_stat(pasta, recursivo=recursivo):
        base = detectar_base_arquivo(item[0])
        if base and base not in por_base:
            por_base[base] = item
    return por_base


def selecionar_arquivos(
    candidatos: Sequence[Candidato],
    modo: ModoLeitura = "ultimo",
) -> list[Candidato]:
    """Seleciona somente o arquivo de maior mtime, sem separar por base."""
    del modo
    if not candidatos:
        return []
    return [max(candidatos, key=lambda item: item[1])]


def contar_arquivos_por_base(candidatos: Sequence[Candidato]) -> dict[str, int]:
    contagem: dict[str, int] = {}
    for caminho, _, _ in candidatos:
        base = detectar_base_arquivo(caminho) or "SEM MARCADOR"
        contagem[base] = contagem.get(base, 0) + 1
    return contagem


# ── Leitura Resiliente (CSV / Excel) ────────────────────────────────
def _ler_bytes_estavel(caminho: str, tamanho_maximo: int) -> bytes:
    try:
        antes = os.stat(caminho)
    except OSError as e:
        raise ArquivoBloqueadoError(f"'{Path(caminho).name}' não pôde ser medido.") from e
    if antes.st_size > tamanho_maximo:
        raise ValueError(
            f"'{Path(caminho).name}' tem {formatar_tamanho(antes.st_size)}, "
            f"acima do limite de {formatar_tamanho(tamanho_maximo)}."
        )
    if antes.st_size == 0:
        raise ValueError(f"'{Path(caminho).name}' está vazio.")
    try:
        with open(caminho, "rb") as fh:
            raw = fh.read()
        depois = os.stat(caminho)
    except PermissionError as e:
        raise ArquivoBloqueadoError(
            f"'{Path(caminho).name}' está bloqueado (OneDrive/Excel)."
        ) from e
    except OSError as e:
        raise ArquivoBloqueadoError(f"'{Path(caminho).name}' ficou inacessível.") from e
    if (
        len(raw) != antes.st_size
        or depois.st_size != antes.st_size
        or depois.st_mtime_ns != antes.st_mtime_ns
        or (depois.st_dev, depois.st_ino) != (antes.st_dev, antes.st_ino)
    ):
        raise ArquivoBloqueadoError(
            f"'{Path(caminho).name}' mudou durante a leitura. Aguardando o próximo ciclo."
        )
    return raw


def _validar_xlsx(raw: bytes, nome: str) -> None:
    if len(raw) < _TAMANHO_MINIMO_XLSX:
        raise ValueError(f"'{nome}' tem {len(raw)} bytes — download incompleto.")
    bio = BytesIO(raw)
    if not zipfile.is_zipfile(bio):
        raise ValueError(f"'{nome}' não é um Excel ZIP válido (zip corrompido ou extensão errada).")
    bio.seek(0)
    with zipfile.ZipFile(bio) as zf:
        ruim = zf.testzip()
    if ruim is not None:
        raise ValueError(f"'{nome}' está corrompido no membro interno: {ruim}")


def _ler_excel_via_openpyxl_nativo(raw: bytes, sheet_name: str | int, header: int) -> pd.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = None
        if isinstance(sheet_name, int):
            if 0 <= sheet_name < len(wb.worksheets):
                ws = wb.worksheets[sheet_name]
        elif sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        if ws is None:
            raise ValueError(f"Aba {sheet_name!r} não encontrada em {wb.sheetnames!r}.")

        linhas = list(ws.iter_rows(values_only=True))
        if header >= len(linhas):
            return pd.DataFrame()
        header_raw = linhas[header]
        if header_raw is None:
            return pd.DataFrame()

        colunas = _nomes_colunas_unicos(header_raw)
        n_cols = len(colunas)
        data: list[list[object]] = []
        for row in linhas[header + 1 :]:
            valores: list[object] = list(row)
            if len(valores) < n_cols:
                valores.extend([None] * (n_cols - len(valores)))
            elif len(valores) > n_cols:
                valores = valores[:n_cols]
            data.append(valores)
        return pd.DataFrame(data, columns=pd.Index(colunas))
    finally:
        wb.close()


def _ler_excel_robusto(
    caminho: str,
    sheet_name: str | int | None = 0,
    header: int = 0,
    tamanho_maximo: int = _TAMANHO_MAXIMO_PADRAO,
) -> tuple[pd.DataFrame, str]:
    planilha: str | int = 0 if sheet_name is None else sheet_name
    ext = Path(caminho).suffix.lower()
    nome = Path(caminho).name
    raw = _ler_bytes_estavel(caminho, tamanho_maximo)
    if ext in (".xlsx", ".xlsm"):
        _validar_xlsx(raw, nome)

    erros: list[str] = []

    def _tentar(
        rotulo: str,
        fn: Callable[[], pd.DataFrame | dict[str, pd.DataFrame]],
    ) -> pd.DataFrame | None:
        try:
            bruto = fn()
            if isinstance(bruto, dict):
                df = next(iter(bruto.values())) if len(bruto) else pd.DataFrame()
            else:
                df = bruto
            if df.empty:
                erros.append(f"{rotulo}: planilha vazia")
                return None
            return df
        except PandasOptionError as e:
            erros.append(f"{rotulo}: OptionError {e}")
            return None
        except ImportError as e:
            erros.append(f"{rotulo}: engine ausente ({e})")
            return None
        except ArquivoBloqueadoError:
            raise
        except PermissionError as e:
            raise ArquivoBloqueadoError(str(e)) from e
        except Exception as e:
            erros.append(f"{rotulo}: {type(e).__name__}: {e}")
            logger.debug("Falha ao ler %s via %s", nome, rotulo, exc_info=True)
            return None

    def _via_temp(engine: str, sufixo: str) -> pd.DataFrame:
        fd, tmp = tempfile.mkstemp(suffix=sufixo)
        os.close(fd)
        try:
            with open(tmp, "wb") as out:
                out.write(raw)
            return pd.read_excel(
                tmp,
                sheet_name=planilha,
                header=header,
                dtype=str,
                engine=cast(EngineExcel, engine),
            )
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass

    def _via_buffer(engine: str, nome_buffer: str) -> pd.DataFrame:
        bio = _BytesNomeados(raw, nome_buffer)
        return pd.read_excel(
            bio,
            sheet_name=planilha,
            header=header,
            dtype=str,
            engine=cast(EngineExcel, engine),
        )

    tentativas: list[tuple[str, Callable[[], pd.DataFrame]]] = []
    if ext == ".xls":
        tentativas.append(("calamine/.xls", lambda: _via_temp("calamine", ".xls")))
        tentativas.append(("xlrd/.xls", lambda: _via_temp("xlrd", ".xls")))
    else:
        tentativas.extend(
            [
                ("openpyxl/temp.xlsx", lambda: _via_temp("openpyxl", ".xlsx")),
                ("openpyxl/BytesIO", lambda: _via_buffer("openpyxl", "arquivo.xlsx")),
                ("calamine", lambda: _via_buffer("calamine", "arquivo.xlsx")),
                (
                    "openpyxl nativo",
                    lambda: _ler_excel_via_openpyxl_nativo(raw, planilha, header),
                ),
            ]
        )

    for rotulo, fn in tentativas:
        df = _tentar(rotulo, fn)
        if df is not None:
            return df, rotulo

    dica = ""
    joined = " | ".join(erros)
    if "engine ausente" in joined.lower():
        dica = " → pip install openpyxl" if ext != ".xls" else " → pip install xlrd"
    raise RuntimeError(f"{joined}{dica}" if erros else f"Falha desconhecida ao ler {nome}.")


_ENCODINGS_CSV: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp1252", "latin1")
_SEPARADORES_CSV: tuple[str, ...] = (";", ",", "\t", "|")


def _quantidade_campos_csv(cabecalho: str, separador: str) -> int:
    """Conta apenas delimitadores fora de aspas (CSV pode contê-los em nomes)."""
    try:
        return len(next(csv.reader([cabecalho], delimiter=separador)))
    except (csv.Error, StopIteration):
        return 0


def _ordem_tentativas_csv(raw: bytes) -> list[tuple[str, str]]:
    """Tenta primeiro o charset e separador mais prováveis, sem parses extras.

    Decodifica a linha de cabeçalho completa (não a corta no meio de um caractere
    UTF-8) e ignora delimitadores dentro de campos entre aspas.
    """
    cabecalho_bytes = raw.split(b"\n", 1)[0]
    primeira_linha = ""
    enc_inicial = _ENCODINGS_CSV[0]
    for enc in _ENCODINGS_CSV:
        try:
            primeira_linha = cabecalho_bytes.decode(enc, "strict")
            enc_inicial = enc
            break
        except UnicodeDecodeError:
            continue

    separadores = sorted(
        _SEPARADORES_CSV,
        key=lambda sep: _quantidade_campos_csv(primeira_linha, sep),
        reverse=True,
    )
    ordem: list[tuple[str, str]] = [(enc_inicial, sep) for sep in separadores]
    ordem.extend((enc, sep) for enc in _ENCODINGS_CSV if enc != enc_inicial for sep in separadores)
    return ordem


def _ler_csv_de_bytes(raw: bytes, nome: str) -> tuple[pd.DataFrame | None, str | None]:
    ultimo: Exception | None = None
    cabecalho_bytes = raw.split(b"\n", 1)[0]
    for enc, sep in _ordem_tentativas_csv(raw):
        try:
            with warnings.catch_warnings():
                # Pandas pode transformar uma coluna excedente em índice ou
                # descartá-la com index_col=False. Nenhum dos dois é aceitável.
                warnings.simplefilter("error", pd.errors.ParserWarning)
                df = pd.read_csv(
                    BytesIO(raw),
                    sep=sep,
                    encoding=enc,
                    dtype=str,
                    low_memory=False,
                    index_col=False,
                    on_bad_lines="error",
                )
        except Exception as e:  # noqa: BLE001 - tentamos o próximo encoding/separador
            ultimo = e
            continue
        if df.shape[1] == 1:
            # Uma coluna é válida; não aceitar um separador errado que deixou
            # intacto um cabeçalho com duas ou mais colunas reais.
            cabecalho = cabecalho_bytes.decode(enc)
            if any(_quantidade_campos_csv(cabecalho, outro) > 1 for outro in _SEPARADORES_CSV):
                continue
        return df, f"csv {enc} sep={sep!r}"
    detalhe = f" ({type(ultimo).__name__}: {ultimo})" if ultimo else ""
    return None, f"{nome}: falha ao interpretar CSV{detalhe}"


def ler_arquivo_detalhado(
    caminho_arquivo: str,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    tamanho_maximo: int = _TAMANHO_MAXIMO_PADRAO,
) -> LeituraArquivo:
    if not caminho_arquivo or not os.path.exists(caminho_arquivo):
        nome = Path(caminho_arquivo).name if caminho_arquivo else "—"
        return LeituraArquivo(None, f"Arquivo não encontrado: {nome}", None)

    nome = Path(caminho_arquivo).name
    ext = Path(caminho_arquivo).suffix.lower()
    if ext not in EXTENSOES_VALIDAS:
        return LeituraArquivo(None, f"{nome}: extensão {ext or '—'} não é CSV/Excel.", None)

    try:
        if ext in (".xlsx", ".xlsm", ".xls"):
            df, motor = _ler_excel_robusto(
                caminho_arquivo,
                sheet_name=sheet_name,
                header=0,
                tamanho_maximo=tamanho_maximo,
            )
            df = _padronizar_dataframe(df)
            if _cabecalho_fraco(list(df.columns)) or not _contem_colunas(df, colunas_esperadas):
                try:
                    alternativo, motor_alt = _ler_excel_robusto(
                        caminho_arquivo,
                        sheet_name=sheet_name,
                        header=1,
                        tamanho_maximo=tamanho_maximo,
                    )
                except Exception as e:  # noqa: BLE001 - header=0 segue válido
                    logger.debug("Header alternativo falhou em %s: %s", nome, e)
                else:
                    alternativo = _padronizar_dataframe(alternativo)
                    alt_melhor = (
                        _contem_colunas(alternativo, colunas_esperadas)
                        and not _contem_colunas(df, colunas_esperadas)
                    ) or (
                        _cabecalho_fraco(list(df.columns))
                        and not _cabecalho_fraco(list(alternativo.columns))
                    )
                    if alt_melhor:
                        df, motor = alternativo, f"{motor_alt} header=1"
        else:
            raw = _ler_bytes_estavel(caminho_arquivo, tamanho_maximo)
            df_csv, motor_csv = _ler_csv_de_bytes(raw, nome)
            if df_csv is None:
                return LeituraArquivo(None, motor_csv, None)
            df, motor = _padronizar_dataframe(df_csv), motor_csv
    except ArquivoBloqueadoError:
        raise
    except Exception as e:
        logger.debug("Falha ao ler %s: %s", nome, e, exc_info=True)
        return LeituraArquivo(None, f"{nome}: {type(e).__name__} — {e}", None)

    if df.empty:
        return LeituraArquivo(None, f"{nome}: arquivo lido, porém sem linhas de dados.", motor)

    df = _alinhar_colunas_esperadas(df, colunas_esperadas)
    faltando = _faltam_colunas(df, colunas_esperadas)
    if faltando:
        return LeituraArquivo(
            None,
            f"{nome}: colunas ausentes → {', '.join(faltando)}",
            motor,
        )
    return LeituraArquivo(df, None, motor)


def processar_candidatos(
    candidatos: Sequence[Candidato],
    modo: ModoLeitura = "ultimo",
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    tamanho_maximo: int = _TAMANHO_MAXIMO_PADRAO,
) -> ResultadoCiclo:
    selecionados = selecionar_arquivos(candidatos, modo)
    if not selecionados:
        return ResultadoCiclo(None, (), (), (), None, (), 0)

    frames: list[pd.DataFrame] = []
    bases: list[str] = []
    arquivos: list[str] = []
    erros: list[str] = []
    motores: list[str] = []
    bloqueios: list[str] = []

    for caminho, _, _ in selecionados:
        try:
            leitura = ler_arquivo_detalhado(
                caminho,
                colunas_esperadas=colunas_esperadas,
                sheet_name=sheet_name,
                tamanho_maximo=tamanho_maximo,
            )
        except ArquivoBloqueadoError as e:
            bloqueios.append(str(e))
            continue
        if leitura.erro or leitura.df is None:
            erros.append(leitura.erro or f"{Path(caminho).name}: leitura vazia")
            continue
        df_parcial = leitura.df
        base = detectar_base_arquivo(caminho)
        if base:
            df_parcial = _aplicar_base_no_dataframe(df_parcial, base)
            if base not in bases:
                bases.append(base)
        frames.append(df_parcial)
        arquivos.append(caminho)
        if leitura.motor:
            motores.append(f"{Path(caminho).name}: {leitura.motor}")

    if not frames:
        return ResultadoCiclo(
            None,
            tuple(bases),
            (),
            tuple(erros),
            bloqueios[0] if bloqueios else None,
            tuple(motores),
            0,
        )

    duplicatas = 0
    if len(frames) == 1:
        df = frames[0]
    else:
        unificados = _unificar_nomes_colunas(frames)
        df = pd.concat(unificados, ignore_index=True, sort=False)
        antes = len(df)
        df = df.drop_duplicates(ignore_index=True)
        duplicatas = antes - len(df)

    return ResultadoCiclo(
        df,
        tuple(bases),
        tuple(arquivos),
        tuple(erros),
        bloqueios[0] if bloqueios else None,
        tuple(motores),
        duplicatas,
    )


# ── Fachada OO (Compatibilidade com TotaleRoboEngine) ───────────────
class TotaleRoboEngine:
    """Fachada unificada de detecção e leitura de exportações TOTALE."""

    PADROES = PADROES_TOTALE
    MAPA_BASES = MAPA_BASES_ARQUIVO

    _ultimo_erro: str | None = None
    _ultimo_motor: str | None = None

    @classmethod
    def ultimo_erro(cls) -> str | None:
        return cls._ultimo_erro

    @classmethod
    def ultimo_motor(cls) -> str | None:
        return cls._ultimo_motor

    @classmethod
    def _registrar(cls, erro: str | None, motor: str | None = None) -> None:
        cls._ultimo_erro = erro
        cls._ultimo_motor = motor
        if erro:
            logger.debug("%s", erro)

    @staticmethod
    def detectar_base(caminho: str | Path | None) -> str | None:
        return detectar_base_arquivo(caminho)

    @staticmethod
    def listar_validos(pasta: str | Path, *, recursivo: bool = False) -> list[Path]:
        candidatos = _listar_candidatos_com_stat(str(pasta), recursivo=recursivo)
        return [Path(c[0]) for c in candidatos]

    @staticmethod
    def buscar_mais_recente(pasta: str | Path) -> Path | None:
        caminho, _, _ = buscar_arquivo_mais_recente(str(pasta))
        return Path(caminho) if caminho else None

    @staticmethod
    def ler_mais_recente(
        pasta: str | Path,
        *,
        aplicar_base: bool = False,
        sheet_name: str | int = 0,
    ) -> pd.DataFrame | None:
        caminho = TotaleRoboEngine.buscar_mais_recente(pasta)
        if caminho is None:
            TotaleRoboEngine._registrar(f"Nenhum arquivo TOTALE válido em {pasta}")
            return None
        return TotaleRoboEngine.ler_dados(
            caminho,
            aplicar_base=aplicar_base,
            sheet_name=sheet_name,
        )

    @staticmethod
    def arquivos_por_base(pasta: str | Path) -> dict[str, Path]:
        """Arquivo mais recente de cada base encontrada na pasta."""
        return {
            base: Path(item[0])
            for base, item in mapear_arquivos_por_base(str(pasta)).items()
        }

    @staticmethod
    def assinatura(caminho: str | Path) -> tuple[int, int] | None:
        try:
            stat = Path(caminho).stat()
        except OSError:
            return None
        return int(stat.st_size), int(stat.st_mtime_ns)

    @staticmethod
    def verificar_estabilidade(
        caminho: str | Path,
        ciclos: int = 2,
        intervalo: float = 0.8,
    ) -> bool:
        arquivo = Path(caminho)
        if not arquivo.is_file():
            return False
        try:
            anterior = arquivo.stat()
        except OSError:
            return False
        if anterior.st_size <= 0:
            return False
        for _ in range(max(ciclos, 1)):
            time.sleep(intervalo)
            try:
                atual = arquivo.stat()
            except OSError:
                return False
            if atual.st_size != anterior.st_size or atual.st_mtime_ns != anterior.st_mtime_ns:
                return False
            anterior = atual
        return anterior.st_size > 0

    @staticmethod
    def ler_dados(
        caminho: str | Path,
        *,
        aplicar_base: bool = False,
        sheet_name: str | int = 0,
    ) -> pd.DataFrame | None:
        try:
            leitura = ler_arquivo_detalhado(
                str(caminho),
                sheet_name=sheet_name,
            )
        except ArquivoBloqueadoError as e:
            TotaleRoboEngine._registrar(str(e))
            return None

        if leitura.erro or leitura.df is None:
            TotaleRoboEngine._registrar(leitura.erro, leitura.motor)
            return None

        df = leitura.df
        if aplicar_base:
            base = TotaleRoboEngine.detectar_base(caminho)
            if base:
                df = _aplicar_base_no_dataframe(df, base)
        TotaleRoboEngine._registrar(None, leitura.motor)
        return df


try:
    import streamlit as st
except ImportError:  # o motor e a CLI também funcionam sem Streamlit
    st = None  # type: ignore[assignment]

_LOCK_STALE_SEGUNDOS: int = 180
EtlFn = Callable[[pd.DataFrame, pd.DataFrame], "pd.DataFrame | None"]
FonteSheets = Callable[[], pd.DataFrame] | pd.DataFrame | None


def ler_arquivo_totale(
    caminho_arquivo: str,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    notificar: bool = True,
) -> pd.DataFrame | None:
    leitura = ler_arquivo_detalhado(caminho_arquivo, colunas_esperadas, sheet_name)
    if st is None or not notificar:
        return leitura.df
    nome = Path(caminho_arquivo).name if caminho_arquivo else "—"
    if leitura.erro:
        st.session_state["robo_erro"] = leitura.erro
        _toast_dedup(f"Falha ao ler: {nome}", icon="⚠️")
        return None
    _toast_dedup(f"Arquivo carregado: {nome}", icon="✅")
    return leitura.df


# ── UI e Controle de Sessão Streamlit ───────────────────────────────
_CSS_ROBO = """
<style>
/* ═══════════════════════════════════════════════════════════════
   ROBÔ DE SINCRONISMO — painel local da Central de Atualização
   Tipografia: títulos em Oscine (--font-hero), dados em Manrope/Inter.
   ═══════════════════════════════════════════════════════════════ */
.totale-robo { margin: 2px 0 10px; }
.totale-robo .robo-panel {
    position: relative; overflow: hidden;
    padding: 26px 28px 22px;
    border: 1px solid #DCE5F0; border-radius: 18px;
    background: linear-gradient(180deg, #FFFFFF 0%, #F8FBFF 100%);
    box-shadow: 0 14px 34px rgba(10, 47, 107, .08);
}
.totale-robo .robo-panel::before {
    content: ''; position: absolute; inset: 0 0 auto 0; height: 4px;
    background: linear-gradient(90deg, #0A2F6B 0%, #2E6DB4 55%, #E8700A 100%);
}
.totale-robo .robo-panel::after {
    content: ''; position: absolute; width: 280px; height: 280px;
    right: -120px; top: -150px; border-radius: 50%;
    background: radial-gradient(circle, rgba(232, 112, 10, .10), transparent 70%);
    pointer-events: none;
}
.totale-robo .robo-head {
    position: relative; z-index: 1;
    display: flex; align-items: flex-start; justify-content: space-between;
    gap: 18px; flex-wrap: wrap;
}
.totale-robo .robo-head-main { display: flex; align-items: flex-start; gap: 15px; min-width: 0; }
.totale-robo .robo-icon {
    display: inline-flex; align-items: center; justify-content: center;
    width: 52px; height: 52px; flex: none; border-radius: 16px;
    color: #FFFFFF;
    background: linear-gradient(135deg, #0A2F6B, #1E5FA8 70%, #2E6DB4);
    box-shadow: 0 10px 22px rgba(10, 47, 107, .22);
}
.totale-robo .robo-icon .material-symbols-rounded { font-size: 27px !important; color: #FFFFFF; }
.totale-robo .robo-eyebrow {
    font-size: 10px; font-weight: 800; letter-spacing: .14em;
    text-transform: uppercase; color: #6B7C93; margin-bottom: 4px;
}
.totale-robo .robo-title {
    font-family: var(--font-hero) !important;
    font-size: clamp(21px, 2.3vw, 27px); font-weight: 800;
    letter-spacing: -.025em; line-height: 1.18;
    color: #0A2F6B; margin: 0;
}
.totale-robo .robo-sub {
    font-size: 12.5px; line-height: 1.55; color: #4B5563;
    margin: 7px 0 0; max-width: 700px;
}
.totale-robo .robo-sub code { font-size: 11.5px; color: #0A2F6B; background: #EEF4FB; padding: 1px 5px; border-radius: 4px; }
.totale-robo .robo-status {
    display: inline-flex; align-items: center; gap: 7px; flex: none;
    padding: 7px 12px; border-radius: 999px;
    font-size: 11px; font-weight: 800; letter-spacing: .02em;
    border: 1px solid transparent; white-space: nowrap;
}
.totale-robo .robo-status--ok   { background: #E5F7EF; color: #087A55; border-color: #B7E7D3; }
.totale-robo .robo-status--wait { background: #FFF4E6; color: #9A4B00; border-color: #F8D8AF; }
.totale-robo .robo-status--off  { background: #EFF2F6; color: #56657A; border-color: #DCE2EA; }
.totale-robo .robo-status--erro { background: #FDEBED; color: #A61B1F; border-color: #F3C3C7; }
.totale-robo .robo-dot {
    width: 7px; height: 7px; border-radius: 50%;
    background: currentColor; box-shadow: 0 0 0 3px rgba(255, 255, 255, .55);
}
.totale-robo .robo-status--ok .robo-dot { animation: robo-pulse 2.4s ease-out infinite; }
@keyframes robo-pulse {
    0%, 100% { box-shadow: 0 0 0 0 rgba(13, 150, 104, .35); }
    50%      { box-shadow: 0 0 0 5px rgba(13, 150, 104, 0); }
}
.totale-robo .robo-stats {
    position: relative; z-index: 1;
    display: grid; grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 10px; margin-top: 20px;
}
.totale-robo .robo-stat {
    background: #FFFFFF; border: 1px solid #E6EDF5;
    border-radius: 12px; padding: 11px 13px; min-width: 0;
}
.totale-robo .robo-stat span {
    display: block; font-size: 9.5px; font-weight: 800;
    letter-spacing: .075em; text-transform: uppercase; color: #6B7C93;
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
}
.totale-robo .robo-stat strong {
    display: block; margin-top: 3px;
    font-family: var(--font-titulo) !important;
    font-size: 21px; font-weight: 800; letter-spacing: -.03em;
    color: #0F172A; font-variant-numeric: tabular-nums;
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
}
.totale-robo .robo-stat small { display: block; margin-top: 1px; font-size: 10px; color: #7C8CA2; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }

/* Controles logo abaixo do painel. */
[class*="st-key-robo_controles_totale"] {
    padding: 11px 15px; margin: 12px 0 6px;
    border: 1px solid #E2EAF4; border-radius: 14px;
    background: #F8FAFC;
}
[class*="st-key-robo_controles_totale"] [data-testid="stWidgetLabel"] p { font-weight: 700; }
[class*="st-key-robo_controles_totale"] .stButton button { min-height: 38px; }

/* Cabeçalhos internos. */
.totale-robo .robo-subhead {
    display: flex; align-items: center; gap: 8px;
    margin: 22px 0 10px;
    font-size: 10.5px; font-weight: 800; letter-spacing: .09em;
    text-transform: uppercase; color: #64748B;
}
.totale-robo .robo-subhead::after { content: ''; flex: 1; height: 1px; background: #E5E7EB; }
.totale-robo .robo-subhead .material-symbols-rounded { font-size: 16px !important; color: #0A2F6B; }
.totale-robo .robo-subhead .robo-count {
    display: inline-flex; align-items: center; justify-content: center;
    min-width: 22px; height: 18px; padding: 0 6px; border-radius: 999px;
    background: #EAF2FC; color: #0A2F6B; font-size: 9.5px; font-weight: 800;
    letter-spacing: 0; text-transform: none;
}

/* Cards das bases. */
.totale-robo .robo-base-card {
    height: 100%; min-height: 158px; box-sizing: border-box;
    display: flex; flex-direction: column;
    padding: 15px 16px 14px;
    background: #FFFFFF; border: 1px solid #E5EAF2;
    border-top: 3px solid #CBD5E1; border-radius: 13px;
    box-shadow: 0 4px 12px rgba(10, 47, 107, .045);
    transition: transform .16s ease, box-shadow .16s ease, border-color .16s ease;
}
.totale-robo .robo-base-card:hover { transform: translateY(-2px); box-shadow: 0 10px 22px rgba(10, 47, 107, .10); }
.totale-robo .robo-base-card--ok   { border-top-color: #0D9668; }
.totale-robo .robo-base-card--wait { border-top-color: #E8700A; }
.totale-robo .robo-base-card--off  { border-top-color: #CBD5E1; background: #FBFCFE; }
.totale-robo .robo-base-top { display: flex; align-items: center; justify-content: space-between; gap: 8px; }
.totale-robo .robo-base-marker {
    display: inline-flex; align-items: center; justify-content: center;
    height: 20px; padding: 0 7px; border-radius: 5px;
    background: #0A2F6B; color: #FFFFFF;
    font-size: 9px; font-weight: 800; letter-spacing: .07em;
}
.totale-robo .robo-base-name {
    font-family: var(--font-hero) !important;
    font-size: 16.5px; font-weight: 800; letter-spacing: -.02em;
    color: #0F172A; margin: 9px 0 0;
}
.totale-robo .robo-base-file {
    margin-top: 5px; font-size: 11.5px; font-weight: 600; color: #334155;
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
}
.totale-robo .robo-base-meta {
    display: flex; align-items: center; justify-content: space-between; gap: 8px;
    margin-top: auto; padding-top: 10px; border-top: 1px dashed #EDF1F6;
    font-size: 10px; color: #7C8CA2; font-variant-numeric: tabular-nums;
}
.totale-robo .robo-chip {
    display: inline-flex; align-items: center; gap: 4px; flex: none;
    padding: 3px 7px; border-radius: 999px;
    font-size: 9px; font-weight: 800; letter-spacing: .03em;
    white-space: nowrap;
}
.totale-robo .robo-chip--ok   { background: #E5F7EF; color: #087A55; }
.totale-robo .robo-chip--wait { background: #FFF0E1; color: #9A4B00; }
.totale-robo .robo-chip--off  { background: #EFF2F6; color: #56657A; }

/* Fila de arquivos. */
.totale-robo .robo-queue {
    border: 1px solid #E6EDF5; border-radius: 14px;
    background: #FFFFFF; overflow: hidden;
}
.totale-robo .robo-queue-row {
    display: flex; align-items: center; gap: 12px;
    padding: 10px 14px; border-bottom: 1px solid #F1F5F9;
}
.totale-robo .robo-queue-row:last-child { border-bottom: 0; }
.totale-robo .robo-queue-icon {
    display: inline-flex; align-items: center; justify-content: center;
    width: 30px; height: 30px; flex: none; border-radius: 9px;
    background: #EEF4FB; color: #0A2F6B; font-size: 17px !important;
}
.totale-robo .robo-queue-main { flex: 1; min-width: 0; }
.totale-robo .robo-queue-name {
    font-size: 12px; font-weight: 650; color: #0F172A;
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
}
.totale-robo .robo-queue-meta { margin-top: 1px; font-size: 10px; color: #7C8CA2; }

/* Estado vazio e passos. */
.totale-robo .robo-empty {
    display: flex; align-items: center; gap: 14px;
    padding: 22px 24px; border: 1px dashed #CBD5E1; border-radius: 14px;
    background: #F8FAFC;
}
.totale-robo .robo-empty .material-symbols-rounded { font-size: 30px !important; color: #94A3B8; }
.totale-robo .robo-empty strong { display: block; font-size: 13px; color: #334155; }
.totale-robo .robo-empty span { display: block; margin-top: 2px; font-size: 11px; color: #7C8CA2; }
.totale-robo .robo-steps {
    display: grid; grid-template-columns: repeat(4, minmax(0, 1fr));
    gap: 8px; margin-top: 16px;
}
.totale-robo .robo-step {
    padding: 10px 12px; border: 1px solid #E9EFF6; border-radius: 11px;
    background: #FFFFFF;
}
.totale-robo .robo-step b {
    display: flex; align-items: center; gap: 6px;
    font-size: 11px; color: #0A2F6B;
}
.totale-robo .robo-step b span {
    display: inline-flex; align-items: center; justify-content: center;
    width: 17px; height: 17px; border-radius: 50%;
    background: #EAF2FC; color: #0A2F6B; font-size: 8.5px; font-weight: 800;
}
.totale-robo .robo-step small { display: block; margin-top: 3px; font-size: 10px; line-height: 1.45; color: #7C8CA2; }

@media (max-width: 900px) {
    .totale-robo .robo-stats { grid-template-columns: repeat(2, minmax(0, 1fr)); }
    .totale-robo .robo-steps { grid-template-columns: repeat(2, minmax(0, 1fr)); }
}
@media (max-width: 640px) {
    .totale-robo .robo-panel { padding: 22px 18px 18px; border-radius: 14px; }
    .totale-robo .robo-stats { grid-template-columns: 1fr; }
    .totale-robo .robo-steps { grid-template-columns: 1fr; }
}
@media (prefers-reduced-motion: reduce) {
    .totale-robo .robo-status--ok .robo-dot { animation: none; }
    .totale-robo .robo-base-card { transition: none; }
    .totale-robo .robo-base-card:hover { transform: none; }
}
</style>
"""


def _exigir_streamlit() -> object:
    if st is None:
        raise RuntimeError("Streamlit não está instalado. A interface do robô só roda no app.")
    return st


def _toast_dedup(mensagem: str, icon: str = "✅") -> None:
    if st is None:
        return
    assinatura = f"{icon}:{mensagem}"
    if st.session_state.get("_robo_ultimo_toast") != assinatura:
        st.session_state["_robo_ultimo_toast"] = assinatura
        st.toast(mensagem, icon=icon)


def _estado_str(chave: str, padrao: str) -> str:
    if st is None:
        return padrao
    valor = st.session_state.get(chave, padrao)
    return str(valor) if valor else padrao


def _estado_bool(chave: str, padrao: bool) -> bool:
    if st is None:
        return padrao
    valor = st.session_state.get(chave, padrao)
    return bool(valor)


def _modo_atual() -> ModoLeitura:
    """O modo é fixo: somente o arquivo de maior mtime entra no pipeline."""
    return "ultimo"


def _lock_ativo() -> bool:
    if st is None or not st.session_state.get("robo_processando", False):
        return False
    inicio = st.session_state.get("robo_processando_desde")
    if isinstance(inicio, datetime):
        agora = datetime.now(tz=inicio.tzinfo)
        idade = (agora - inicio).total_seconds()
        if 0 <= idade < _LOCK_STALE_SEGUNDOS:
            return True
    st.session_state["robo_processando"] = False
    logger.warning("Lock do robô expirou ou não tem horário de início válido.")
    return False


def _obter_gsheets(gsheets_fn: FonteSheets) -> pd.DataFrame:
    if callable(gsheets_fn):
        resultado = gsheets_fn()
        return resultado if isinstance(resultado, pd.DataFrame) else pd.DataFrame()
    if isinstance(gsheets_fn, pd.DataFrame):
        return gsheets_fn
    return pd.DataFrame()


def _rerun_aplicacao() -> None:
    if st is None:
        return
    rerun = st.rerun
    try:
        rerun(scope="app")
    except TypeError:
        rerun()


def _data_curta(mtime_ns: int) -> str:
    try:
        return datetime.fromtimestamp(mtime_ns / 1_000_000_000).strftime("%d/%m %H:%M")
    except (OverflowError, OSError, ValueError):
        return "—"


def _rotulo_status(
    ativo: bool,
    sig_ok: bool,
    erro: str | None,
    aguardando: str | None,
) -> tuple[str, str]:
    if not ativo:
        return "Pausado", "off"
    if erro:
        return "Atenção", "erro"
    if sig_ok:
        return "Sincronizado", "ok"
    if aguardando:
        return "Aguardando arquivo", "wait"
    return "Monitorando", "wait"


def _html_painel_robo(
    *,
    pasta: str,
    ativo: bool,
    sig_ok: bool,
    erro: str | None,
    aguardando: str | None,
    n_candidatos: int,
    n_fila: int,
    n_bases: int,
    hora_sucesso: str,
    ultimo_nome: str,
) -> str:
    status_texto, status_classe = _rotulo_status(ativo, sig_ok, erro, aguardando)
    pasta_nome = (Path(pasta).name if pasta else "") or "pasta não definida"
    if not pasta.strip():
        resumo = "Defina uma pasta monitorada em <b>⚙️ Pasta monitorada</b> para começar."
    elif not ativo:
        resumo = (
            "A sincronização automática está pausada. Ative o robô para acompanhar "
            "novas exportações nesta pasta."
        )
    elif n_candidatos == 0:
        resumo = (
            "Nenhuma exportação TOTALE encontrada. Salve o arquivo em "
            f"<code>{escape(pasta)}</code> e o robô assume o restante."
        )
    elif sig_ok:
        resumo = (
            "A base mais recente já foi lida, tratada e publicada para os painéis. "
            "Novos arquivos entram na fila automaticamente."
        )
    else:
        resumo = (
            f"Acompanhando <code>{escape(pasta)}</code> a cada 5 segundos. "
            "O arquivo mais recente entra no pipeline assim que ficar estável."
        )
    return (
        '<div class="totale-robo"><div class="robo-panel">'
        '<div class="robo-head">'
        '<div class="robo-head-main">'
        '<div class="robo-icon"><span class="material-symbols-rounded" aria-hidden="true">'
        "smart_toy</span></div>"
        '<div><div class="robo-eyebrow">Robô de sincronismo</div>'
        '<h2 class="robo-title">Bases locais prontas para os painéis</h2>'
        f'<p class="robo-sub">{resumo}</p></div>'
        "</div>"
        f'<div class="robo-status robo-status--{status_classe}">'
        f'<span class="robo-dot" aria-hidden="true"></span>{escape(status_texto)}</div>'
        "</div>"
        '<div class="robo-stats">'
        f'<div class="robo-stat"><span>Pasta monitorada</span><strong>{escape(pasta_nome)}</strong>'
        f"<small>{escape(pasta)}</small></div>"
        f'<div class="robo-stat"><span>Arquivos na pasta</span><strong>{n_candidatos}</strong>'
        f'<small>{n_fila} no pipeline agora</small></div>'
        f'<div class="robo-stat"><span>Bases detectadas</span><strong>{n_bases}/3</strong>'
        '<small>SPO · GRS · ABCDM</small></div>'
        f'<div class="robo-stat"><span>Última sincronização</span><strong>{escape(hora_sucesso)}</strong>'
        f"<small>{escape(ultimo_nome or 'Aguardando a primeira leitura')}</small></div>"
        "</div></div></div>"
    )


def _html_card_base(
    marcador: str,
    base: str,
    info: Candidato | None,
    processados_map: dict[object, object],
    selecionados: Sequence[Candidato],
) -> str:
    if info is None:
        return (
            '<div class="totale-robo"><div class="robo-base-card robo-base-card--off">'
            '<div class="robo-base-top">'
            f'<span class="robo-base-marker">{escape(marcador)}</span>'
            '<span class="robo-chip robo-chip--off">Aguardando</span></div>'
            f'<h3 class="robo-base-name">{escape(base)}</h3>'
            f'<div class="robo-base-file">Nenhum arquivo com “{escape(marcador)}” no nome</div>'
            '<div class="robo-base-meta"><span>Sem exportação</span><span>—</span></div>'
            "</div></div>"
        )

    caminho, mtime_ns, size = info
    meta = obter_metadados_arquivo(caminho)
    sincronizado = processados_map.get(caminho) == mtime_ns
    na_fila = any(item[0] == caminho for item in selecionados)
    if sincronizado:
        classe, chip, rotulo = "ok", "robo-chip--ok", "Sincronizada"
    elif na_fila:
        classe, chip, rotulo = "wait", "robo-chip--wait", "Na fila"
    else:
        classe, chip, rotulo = "off", "robo-chip--off", "Mais antigo"
    return (
        f'<div class="totale-robo"><div class="robo-base-card robo-base-card--{classe}">'
        '<div class="robo-base-top">'
        f'<span class="robo-base-marker">{escape(marcador)}</span>'
        f'<span class="robo-chip {chip}">{rotulo}</span></div>'
        f'<h3 class="robo-base-name">{escape(base)}</h3>'
        f'<div class="robo-base-file" title="{escape(caminho, quote=True)}">'
        f"{escape(meta['nome'])}</div>"
        '<div class="robo-base-meta">'
        f"<span>{escape(formatar_tamanho(size))}</span>"
        f"<span>{escape(_data_curta(mtime_ns))}</span>"
        "</div></div></div>"
    )


def _html_fila_arquivos(
    candidatos: Sequence[Candidato],
    selecionados: Sequence[Candidato],
    processados_map: dict[object, object],
    limite: int = 5,
) -> str:
    selecionados_paths = {item[0] for item in selecionados}
    linhas: list[str] = []
    for caminho, mtime_ns, size in candidatos[:limite]:
        base = detectar_base_arquivo(caminho) or "sem marcador"
        sincronizado = processados_map.get(caminho) == mtime_ns
        if sincronizado:
            icone, chip, rotulo = "check_circle", "robo-chip--ok", "Sincronizado"
        elif caminho in selecionados_paths:
            icone, chip, rotulo = "bolt", "robo-chip--wait", "Na fila"
        else:
            icone, chip, rotulo = "history", "robo-chip--off", "Fora do pipeline"
        linhas.append(
            '<div class="robo-queue-row">'
            f'<span class="robo-queue-icon material-symbols-rounded" aria-hidden="true">{icone}</span>'
            '<div class="robo-queue-main">'
            f'<div class="robo-queue-name" title="{escape(caminho, quote=True)}">'
            f"{escape(Path(caminho).name)}</div>"
            f'<div class="robo-queue-meta">{escape(base)} · {escape(formatar_tamanho(size))} '
            f"· {escape(_data_curta(mtime_ns))}</div></div>"
            f'<span class="robo-chip {chip}">{rotulo}</span></div>'
        )
    restantes = len(candidatos) - limite
    if restantes > 0:
        linhas.append(
            '<div class="robo-queue-row">'
            '<span class="robo-queue-icon material-symbols-rounded" aria-hidden="true">more_horiz</span>'
            f'<div class="robo-queue-main"><div class="robo-queue-name">+{restantes} arquivo(s) '
            "mais antigo(s)</div><div class=\"robo-queue-meta\">Mantidos apenas como histórico "
            "da pasta</div></div>"
            '<span class="robo-chip robo-chip--off">Histórico</span></div>'
        )
    return '<div class="totale-robo"><div class="robo-queue">' + "".join(linhas) + "</div></div>"


def _html_pasta_vazia(pasta: str, ativo: bool) -> str:
    if ativo:
        titulo = "Monitorando novas exportações"
        detalhe = (
            "Nenhum arquivo TOTALE nesta pasta ainda. Salve a exportação em "
            f"{pasta} — o robô detecta sozinho."
        )
    else:
        titulo = "Sincronização pausada"
        detalhe = "Ative o auto-sincronizar para retomar a leitura da pasta."
    return (
        '<div class="totale-robo"><div class="robo-empty">'
        '<span class="material-symbols-rounded" aria-hidden="true">folder_open</span>'
        f"<div><strong>{escape(titulo)}</strong><span>{escape(detalhe)}</span></div>"
        "</div></div>"
    )


def _html_passos_robo() -> str:
    passos = (
        ("Detecta", "Varre a pasta e ignora temporários, atalhos e downloads incompletos."),
        ("Valida", "Confere estabilidade, tamanho, ZIP do Excel e encoding do CSV."),
        ("Trata", "Padroniza texto, bases e colunas antes de qualquer publicação."),
        ("Publica", "Atualiza a base da sessão e libera os painéis sem upload manual."),
    )
    itens = "".join(
        f'<div class="robo-step"><b><span>{i:02d}</span>{escape(titulo)}</b>'
        f"<small>{escape(detalhe)}</small></div>"
        for i, (titulo, detalhe) in enumerate(passos, start=1)
    )
    return f'<div class="totale-robo"><div class="robo-steps">{itens}</div></div>'


def _executar_verificacao_robo(
    pasta_monitorada: str,
    etl_fn: EtlFn | None = None,
    gsheets_fn: FonteSheets = None,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    ciclos_estabilidade: int = 2,
    modo: ModoLeitura = "ultimo",
    recursivo: bool = False,
) -> None:
    if st is None or not _estado_bool("robo_ativo", True) or _lock_ativo():
        return

    candidatos = _listar_candidatos_com_stat(pasta_monitorada, recursivo=recursivo)
    selecionados = selecionar_arquivos(candidatos, modo)
    if not selecionados:
        return

    assinatura = tuple(selecionados)
    if st.session_state.get("robo_candidato_sig") == assinatura:
        stable = int(st.session_state.get("robo_candidato_stable", 0) or 0) + 1
    else:
        stable = 1
    st.session_state["robo_candidato_sig"] = assinatura
    st.session_state["robo_candidato_stable"] = stable
    if stable < max(ciclos_estabilidade, 1):
        return
    if st.session_state.get("robo_processado_sig") == assinatura:
        return
    # Um arquivo corrompido ou um ETL quebrado não deve ser reprocessado a cada
    # 5 s. A falha fica presa à assinatura; um arquivo novo (ou "Tentar
    # novamente") libera a próxima tentativa.
    if st.session_state.get("robo_falha_sig") == assinatura:
        return

    st.session_state["robo_processando"] = True
    st.session_state["robo_processando_desde"] = datetime.now()
    try:
        if (
            tuple(
                selecionar_arquivos(
                    _listar_candidatos_com_stat(pasta_monitorada, recursivo=recursivo),
                    modo,
                )
            )
            != assinatura
        ):
            st.session_state["robo_candidato_stable"] = 0
            _toast_dedup("Arquivos mudaram durante a espera. Aguardando...", icon="⚠️")
            return

        with st.spinner("🤖 Lendo a exportação mais recente..."):
            resultado = processar_candidatos(
                candidatos,
                modo=modo,
                colunas_esperadas=colunas_esperadas,
                sheet_name=sheet_name,
            )
        if resultado.bloqueio and resultado.df is None:
            st.session_state["robo_candidato_stable"] = 0
            st.session_state["_robo_aguardando_lock"] = resultado.bloqueio
            return
        if resultado.df is None:
            st.session_state.pop("_robo_aguardando_lock", None)
            if resultado.erros:
                st.session_state["robo_erro"] = " | ".join(resultado.erros)
                st.session_state["robo_falha_sig"] = assinatura
            return

        def _assinatura_atual() -> tuple[Candidato, ...]:
            atuais = _listar_candidatos_com_stat(pasta_monitorada, recursivo=recursivo)
            return tuple(selecionar_arquivos(atuais, modo))

        if _assinatura_atual() != assinatura:
            st.session_state["robo_candidato_stable"] = 0
            _toast_dedup("Arquivos mudaram durante a leitura. Aguardando...", icon="⚠️")
            return

        # Executa o ETL antes de publicar a nova assinatura e os dados. Uma falha
        # não deve marcar o arquivo como processado nem substituir a base anterior.
        bruto = resultado.df
        processado = bruto
        if callable(etl_fn):
            with st.spinner("🧩 Tratando a base para os painéis..."):
                transformado = etl_fn(bruto.copy(deep=True), _obter_gsheets(gsheets_fn))
            if transformado is not None:
                if not isinstance(transformado, pd.DataFrame):
                    raise TypeError("O ETL do robô deve retornar DataFrame ou None.")
                processado = transformado
        if _assinatura_atual() != assinatura:
            st.session_state["robo_candidato_stable"] = 0
            _toast_dedup("Arquivos mudaram durante o ETL. Aguardando...", icon="⚠️")
            return

        caminho_recente = selecionados[0][0]
        st.session_state["df_memoria_raw"] = bruto
        st.session_state["df_memoria"] = processado
        st.session_state["robo_ultimo_processado_path"] = caminho_recente
        st.session_state["robo_ultimo_processado_mtime"] = selecionados[0][1]
        st.session_state["robo_arquivos_processados"] = {
            caminho: mtime for caminho, mtime, _ in selecionados
        }
        st.session_state["robo_bases_carregadas"] = list(resultado.bases)
        st.session_state["robo_base_detectada"] = detectar_base_arquivo(caminho_recente)
        st.session_state["robo_hora_sucesso"] = datetime.now()
        st.session_state["robo_motores"] = list(resultado.motores)
        st.session_state["robo_duplicatas_removidas"] = resultado.duplicatas_removidas
        st.session_state.pop("_robo_aguardando_lock", None)
        st.session_state.pop("robo_falha_sig", None)

        sufixo = f" · {', '.join(resultado.bases)}" if resultado.bases else ""
        st.session_state["origem_dados"] = (
            f"Robô Local ({len(resultado.arquivos)} arquivo(s){sufixo})"
        )
        if resultado.erros:
            st.session_state["robo_erro"] = "Parcial: " + " | ".join(resultado.erros)
        else:
            st.session_state.pop("robo_erro", None)
        st.session_state["robo_processado_sig"] = assinatura

        if resultado.bases:
            _toast_dedup(f"Bases: {', '.join(resultado.bases)} — pipeline concluído!", icon="🚀")
        else:
            _toast_dedup("Pipeline concluído!", icon="🚀")
        _rerun_aplicacao()
    except Exception as e:
        logger.exception("Falha no ciclo do robô")
        st.session_state["robo_erro"] = f"{type(e).__name__}: {e}"
        st.session_state["robo_falha_sig"] = assinatura
        _toast_dedup(f"Erro no pipeline: {e}", icon="❌")
    finally:
        st.session_state["robo_processando"] = False
        st.session_state["robo_processando_desde"] = None


if st is not None and hasattr(st, "fragment"):
    _executar_verificacao_robo = st.fragment(run_every="5s")(_executar_verificacao_robo)


def renderizar_robo_local(
    etl_fn: EtlFn | None = None,
    gsheets_fn: FonteSheets = None,
    pasta_padrao: str | None = None,
    colunas_esperadas: list[str] | None = None,
    sheet_name: str | int | None = 0,
    ciclos_estabilidade: int = 2,
    mostrar_toggle: bool = True,
    mostrar_config: bool = True,
    mostrar_diagnostico: bool = True,
) -> None:
    """Renderiza o painel do robô no contêiner atual da Central de Atualização."""
    _exigir_streamlit()
    assert st is not None

    st.markdown(_CSS_ROBO, unsafe_allow_html=True)

    if "robo_ativo" not in st.session_state:
        st.session_state["robo_ativo"] = True
    st.session_state["robo_modo_leitura"] = "ultimo"
    if "robo_recursivo" not in st.session_state:
        st.session_state["robo_recursivo"] = False

    pasta_default = pasta_padrao or obter_pasta_robo_padrao()
    if not st.session_state.get("robo_pasta_alvo"):
        st.session_state["robo_pasta_alvo"] = pasta_default
    elif (
        pasta_padrao
        and not st.session_state.get("_robo_pasta_user_set")
        and st.session_state.get("robo_pasta_alvo") == obter_pasta_robo_padrao()
    ):
        st.session_state["robo_pasta_alvo"] = pasta_padrao

    pasta_alvo = str(st.session_state["robo_pasta_alvo"])
    recursivo = _estado_bool("robo_recursivo", False)
    modo = _modo_atual()
    candidatos_ui = _listar_candidatos_com_stat(pasta_alvo, recursivo=recursivo)
    selecionados_ui = selecionar_arquivos(candidatos_ui, modo)
    por_base = mapear_arquivos_por_base(pasta_alvo, recursivo=recursivo)

    processados_bruto = st.session_state.get("robo_arquivos_processados")
    processados_map: dict[object, object] = (
        processados_bruto if isinstance(processados_bruto, dict) else {}
    )
    erro_bruto = st.session_state.get("robo_erro")
    erro = erro_bruto if isinstance(erro_bruto, str) and erro_bruto else None
    aguardando_bruto = st.session_state.get("_robo_aguardando_lock")
    aguardando = aguardando_bruto if isinstance(aguardando_bruto, str) and aguardando_bruto else None
    assinatura = tuple(selecionados_ui)
    sig_ok = bool(assinatura) and st.session_state.get("robo_processado_sig") == assinatura

    hora_bruta = st.session_state.get("robo_hora_sucesso")
    hora_sucesso = hora_bruta.strftime("%H:%M:%S") if isinstance(hora_bruta, datetime) else "—"
    ultimo_path = st.session_state.get("robo_ultimo_processado_path")
    ultimo_nome = Path(str(ultimo_path)).name if ultimo_path else ""

    cabecalho = st.empty()

    with st.container(key="robo_controles_totale"):
        col_toggle, col_acao, col_info = st.columns([1.35, 0.95, 1.7])
        with col_toggle:
            if mostrar_toggle:
                ativo = st.toggle(
                    "Auto-sincronizar",
                    value=_estado_bool("robo_ativo", True),
                    key="robo_toggle_ui",
                    help=(
                        "Monitora a pasta a cada 5 s enquanto esta sessão estiver aberta. "
                        "SPO→Leste, GRS→Guarulhos, ABCDM→ABCDM."
                    ),
                )
            else:
                ativo = _estado_bool("robo_ativo", True)
            st.session_state["robo_ativo"] = bool(ativo)
        with col_acao:
            sincronizar = st.button(
                "Sincronizar agora",
                key="robo_sincronizar_agora",
                width="stretch",
                disabled=not bool(ativo),
                help=(
                    "Força uma leitura imediata do arquivo mais recente."
                    if ativo
                    else "Ative o auto-sincronizar para liberar a leitura manual."
                ),
            )
        with col_info:
            if candidatos_ui:
                st.caption(
                    f"Leitura automática a cada 5 s · {len(candidatos_ui)} arquivo(s) "
                    f"na pasta · {len(selecionados_ui)} no pipeline"
                )
            else:
                st.caption(
                    "Leitura automática a cada 5 s · aguardando a primeira exportação na pasta"
                )

    if sincronizar:
        if not selecionados_ui:
            _toast_dedup("Nenhum arquivo TOTALE encontrado na pasta monitorada.", icon="⚠️")
        else:
            st.session_state["robo_candidato_sig"] = assinatura
            st.session_state["robo_candidato_stable"] = max(ciclos_estabilidade, 1)
            st.session_state.pop("robo_processado_sig", None)
            st.session_state.pop("robo_falha_sig", None)
            _rerun_aplicacao()

    if mostrar_config:
        with st.expander("⚙️ Pasta monitorada", expanded=False):
            pasta_input = st.text_input(
                "Caminho do disco",
                value=pasta_alvo,
                key="input_pasta_robo",
                help="A pasta em que o TOTALE salva as exportações (normalmente Downloads).",
            )
            if pasta_input != pasta_alvo:
                st.session_state["_robo_pasta_user_set"] = True
                st.session_state["robo_pasta_alvo"] = pasta_input
                _rerun_aplicacao()

            novo_recursivo = st.checkbox(
                "Incluir subpastas",
                value=_estado_bool("robo_recursivo", False),
                key="robo_recursivo_ui",
            )
            if bool(novo_recursivo) != recursivo:
                st.session_state["robo_recursivo"] = bool(novo_recursivo)
                _rerun_aplicacao()

            if os.path.isdir(pasta_alvo):
                st.success("Pasta válida")
            else:
                st.error("Pasta não encontrada — confira o caminho digitado.")
            st.caption(
                "Aceita somente arquivos iniciados com Atividades- (*.csv, .xlsx, .xlsm, .xls). "
                "Temporários (~$, .tmp, .crdownload) são ignorados."
            )

    cabecalho.markdown(
        _html_painel_robo(
            pasta=pasta_alvo,
            ativo=bool(ativo),
            sig_ok=sig_ok,
            erro=erro,
            aguardando=aguardando,
            n_candidatos=len(candidatos_ui),
            n_fila=len(selecionados_ui),
            n_bases=len(por_base),
            hora_sucesso=hora_sucesso,
            ultimo_nome=ultimo_nome,
        ),
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="totale-robo"><div class="robo-subhead">'
        '<span class="material-symbols-rounded" aria-hidden="true">dns</span>'
        f"Bases monitoradas<span class=\"robo-count\">{len(por_base)}/3</span></div></div>",
        unsafe_allow_html=True,
    )
    colunas_bases = st.columns(3)
    for coluna, (marcador, base) in zip(colunas_bases, MAPA_BASES_ARQUIVO.items()):
        with coluna:
            st.markdown(
                _html_card_base(
                    marcador,
                    base,
                    por_base.get(base),
                    processados_map,
                    selecionados_ui,
                ),
                unsafe_allow_html=True,
            )

    if candidatos_ui:
        st.markdown(
            '<div class="totale-robo"><div class="robo-subhead">'
            '<span class="material-symbols-rounded" aria-hidden="true">list_alt</span>'
            "Fila de arquivos"
            f'<span class="robo-count">{len(candidatos_ui)}</span></div></div>',
            unsafe_allow_html=True,
        )
        st.markdown(
            _html_fila_arquivos(candidatos_ui, selecionados_ui, processados_map),
            unsafe_allow_html=True,
        )
    else:
        st.markdown(_html_pasta_vazia(pasta_alvo, bool(ativo)), unsafe_allow_html=True)

    if mostrar_diagnostico:
        with st.expander("🔎 Como o robô trabalha", expanded=False):
            st.markdown(_html_passos_robo(), unsafe_allow_html=True)
            st.divider()
            motores = st.session_state.get("robo_motores")
            if isinstance(motores, list) and motores:
                st.caption("Motor usado na última leitura:")
                for motor in motores:
                    st.caption(f"• {motor}")
            else:
                st.caption("Ainda não houve leitura concluída nesta sessão.")
            contagem = contar_arquivos_por_base(candidatos_ui)
            if contagem:
                detalhe = " · ".join(f"{base}: {n}" for base, n in contagem.items())
                st.caption(f"Arquivos por base na pasta → {detalhe}")
            removidas = st.session_state.get("robo_duplicatas_removidas")
            if isinstance(removidas, int) and removidas > 0:
                st.caption(f"{removidas} linha(s) idêntica(s) removida(s) na junção.")
            if aguardando:
                st.info(f"⏳ {aguardando}")
            st.caption(
                "O pipeline usa somente o arquivo de maior mtime da pasta; os demais "
                "ficam disponíveis como histórico. A base tratada é publicada em "
                "df_memoria para os painéis."
            )

    if erro:
        st.error(f"❌ {erro}")
        col_limpar, col_retry = st.columns(2)
        with col_limpar:
            if st.button("🧹 Limpar", key="robo_limpar_erro", width="stretch"):
                st.session_state.pop("robo_erro", None)
                _rerun_aplicacao()
        with col_retry:
            if st.button("🔄 Tentar novamente", key="robo_retry_erro", width="stretch"):
                st.session_state.pop("robo_erro", None)
                st.session_state.pop("robo_falha_sig", None)
                st.session_state["robo_candidato_sig"] = assinatura
                st.session_state["robo_processado_sig"] = None
                st.session_state["robo_candidato_stable"] = max(ciclos_estabilidade, 1)
                _rerun_aplicacao()

    _executar_verificacao_robo(
        pasta_monitorada=pasta_alvo,
        etl_fn=etl_fn,
        gsheets_fn=gsheets_fn,
        colunas_esperadas=colunas_esperadas,
        sheet_name=sheet_name,
        ciclos_estabilidade=ciclos_estabilidade,
        modo=modo,
        recursivo=recursivo,
    )


def executar_ciclo_headless(
    pasta: str,
    destino_csv: str | None = None,
    modo: ModoLeitura = "ultimo",
    colunas_esperadas: list[str] | None = None,
    recursivo: bool = False,
) -> ResultadoCiclo:
    """Um ciclo sem Streamlit. Processa apenas o arquivo mais recente da pasta."""
    resultado = processar_candidatos(
        _listar_candidatos_com_stat(pasta, recursivo=recursivo),
        modo=modo,
        colunas_esperadas=colunas_esperadas,
    )
    if resultado.df is not None and destino_csv:
        destino = Path(destino_csv)
        if destino.resolve() == Path(resultado.arquivos[0]).resolve():
            raise ValueError("O CSV de saída não pode sobrescrever o arquivo de origem.")
        destino.parent.mkdir(parents=True, exist_ok=True)
        resultado.df.to_csv(destino, index=False, sep=";", encoding="utf-8-sig")
    return resultado


renderizar_sidebar_robo = renderizar_robo_local
_obter_pasta_robo_padrao = obter_pasta_robo_padrao
_obter_metadados_arquivo = obter_metadados_arquivo

__all__ = [
    "EXTENSOES_TEMP",
    "EXTENSOES_VALIDAS",
    "MAPA_BASES_ARQUIVO",
    "PADROES_TOTALE",
    "_TAMANHO_MAXIMO_PADRAO",
    "_TAMANHO_MINIMO_XLSX",
    "_VALORES_AUSENTES",
    "ArquivoBloqueadoError",
    "Candidato",
    "EngineExcel",
    "LeituraArquivo",
    "ModoLeitura",
    "ResultadoCiclo",
    "TotaleRoboEngine",
    "_aplicar_base_no_dataframe",
    "_cabecalho_fraco",
    "_celula_para_texto",
    "_listar_candidatos_com_stat",
    "_nomes_colunas_unicos",
    "_ordem_tentativas_csv",
    "_padronizar_coluna_texto",
    "_padronizar_dataframe",
    "_unificar_nomes_colunas",
    "buscar_arquivo_mais_recente",
    "contar_arquivos_por_base",
    "detectar_base_arquivo",
    "detectar_marcador_arquivo",
    "executar_ciclo_headless",
    "formatar_tamanho",
    "ler_arquivo_detalhado",
    "ler_arquivo_totale",
    "mapear_arquivos_por_base",
    "obter_metadados_arquivo",
    "obter_pasta_robo_padrao",
    "processar_candidatos",
    "renderizar_robo_local",
    "renderizar_sidebar_robo",
    "selecionar_arquivos",
]


if __name__ == "__main__":
    pasta_cli = sys.argv[1] if len(sys.argv) > 1 else obter_pasta_robo_padrao()
    destino = sys.argv[2] if len(sys.argv) > 2 else str(Path(pasta_cli) / "_robo_cache.csv")
    try:
        ciclo = executar_ciclo_headless(pasta_cli, destino_csv=destino)
    except (OSError, ValueError) as exc:
        print(f"Erro ao executar o robô: {exc}", file=sys.stderr)
        raise SystemExit(2) from None
    if ciclo.df is None:
        print(ciclo.bloqueio or " | ".join(ciclo.erros) or "Nenhum arquivo processado.")
        raise SystemExit(1)
    print(f"{len(ciclo.df)} linhas, bases={', '.join(ciclo.bases) or '—'} → {destino}")
